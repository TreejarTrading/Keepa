#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Interactive LEAN UAE launch model: toggle SKUs (Вкл=1/0), budget & payback recompute."""
import json, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.workbook.properties import CalcProperties

BASE="/home/user/Keepa"
d={x['asin']:x for x in json.load(open(f"{BASE}/analysis_workspace/03_unit_economics/uae_final20.json"))}
meta=json.load(open(f"{BASE}/analysis_workspace/04_alibaba/sku_meta.json"))
# heroes: (asin, pilot qty, electrical?, include-in-lean-default)
heroes=[
 ("B0GYF91BXB",300,0,1),("B0GDTFFKYC",300,0,1),("B0GSZHLZ3B",250,0,1),
 ("B0GCJQ4Z51",200,0,1),("B0GZZ3VM69",150,0,1),
 ("B0GKMVTD6V",300,1,0),("B0GV2DHXWS",250,1,0),("B0GF7LLCL8",200,0,0),
 ("B0GL357H51",100,0,0),("B0GQHGFFHC",120,1,0),
]
wb=Workbook()
HDR=PatternFill("solid",fgColor="1F4E78"); SUB=PatternFill("solid",fgColor="2E75B6")
PARAM=PatternFill("solid",fgColor="FFF2CC"); INF=PatternFill("solid",fgColor="E2EFDA")
TOT=PatternFill("solid",fgColor="D9E1F2")
WHITE=Font(color="FFFFFF",bold=True); BOLD=Font(bold=True); LINK=Font(color="0563C1",underline="single")
thin=Side(style="thin",color="BFBFBF"); BORD=Border(left=thin,right=thin,top=thin,bottom=thin)
CEN=Alignment("center","center",wrap_text=True); LEFT=Alignment("left","center",wrap_text=True)

# Params
p=wb.active; p.title="Параметры_лин"; p.sheet_view.showGridLines=False
p["A1"]="ЛИН-ЗАПУСК ОАЭ — редактируйте жёлтые ячейки. На листе 'Лин-план' ставьте Вкл=1/0 по каждому SKU."
p["A1"].font=Font(bold=True,size=12,color="1F4E78"); p.merge_cells("A1:D1")
params=[("USD_AED","Курс USD→AED",3.6725,"0.0000"),
 ("PHOTO","Фото/листинг на SKU, AED",700,"#,##0"),
 ("SAMPLES","Образцы+QC на SKU, AED",500,"#,##0"),
 ("PPC_SKU","Стартовый PPC на SKU, AED",1500,"#,##0"),
 ("LOGI","Логистика (общая), AED",2500,"#,##0"),
 ("CERT","Сертификация (если есть электрика), AED",6000,"#,##0"),
 ("RESERVE","Резерв",0.05,"0%"),
 ("CIT","Налог ОАЭ",0.09,"0.0%")]
p.cell(2,1,"Параметр").font=BOLD; p.cell(2,2,"Значение").font=BOLD
for i,(nm,lb,val,fmt) in enumerate(params):
    r=3+i; p.cell(r,1,lb).alignment=LEFT
    c=p.cell(r,2,val); c.fill=PARAM; c.number_format=fmt; c.font=BOLD; c.border=BORD; c.alignment=CEN
    wb.defined_names[nm]=DefinedName(nm,attr_text=f"'Параметры_лин'!$B${r}")
p.column_dimensions["A"].width=40; p.column_dimensions["B"].width=12
notes=["ИДЕЯ ЛИН-ЗАПУСКА:","• Тестируем нишу малой кровью: несколько SKU-героев, партии близко к MOQ.",
 "• Старт с НЕэлектрических товаров → без сертификации ESMA/G-Mark (быстрее и дешевле).",
 "• PPC на старте скромный, наращиваем из выручки, а не из капитала.",
 "• Ставьте Вкл=1 у нужных SKU — бюджет, EBITDA и окупаемость пересчитаются.",
 "• Прибыль реинвестируем в следующие партии/SKU (самофинансируемый рост)."]
for i,t in enumerate(notes):
    cc=p.cell(12+i,1,t); cc.alignment=LEFT
    if t.endswith(":"): cc.font=BOLD
    p.merge_cells(start_row=12+i,start_column=1,end_row=12+i,end_column=6)

# Lean plan
m=wb.create_sheet("Лин-план"); m.sheet_view.showGridLines=False
m.cell(1,1,"ЛИН-ПЛАН ЗАПУСКА ОАЭ — Вкл(1/0), Заказ шт — ввод; остальное формулы.").font=Font(bold=True,size=12,color="1F4E78")
cols=[("№",4),("Товар",32),("ASIN",13),("Электро?",8),("Вкл (1/0)",8),("Заказ, шт",9),
 ("Landed AED/шт",11),("Закупка AED",11),("EBITDA/шт AED",11),("EBITDA/мес AED",12)]
m.merge_cells(start_row=1,start_column=1,end_row=1,end_column=len(cols))
for j,(n,w) in enumerate(cols,1):
    c=m.cell(3,j,n); c.fill=HDR; c.font=WHITE; c.alignment=CEN; c.border=BORD
    m.column_dimensions[get_column_letter(j)].width=w
first=4
for i,(a,qty,elec,inc) in enumerate(heroes):
    r=first+i; x=d[a]
    m.cell(r,1,i+1).alignment=CEN
    m.cell(r,2,meta[a]["ru"]).alignment=LEFT
    ac=m.cell(r,3,a); ac.hyperlink=f"https://www.amazon.ae/dp/{a}"; ac.font=LINK; ac.alignment=CEN
    m.cell(r,4,"да" if elec else "нет").alignment=CEN
    ve=m.cell(r,5,inc); ve.fill=PARAM; ve.alignment=CEN; ve.border=BORD
    vq=m.cell(r,6,qty); vq.fill=INF; vq.alignment=CEN; vq.border=BORD; vq.number_format="#,##0"
    m.cell(r,7,round(x["s_landed_aed"],1)).number_format="#,##0.0"
    m.cell(r,7).fill=INF; m.cell(r,7).alignment=CEN; m.cell(r,7).border=BORD
    m.cell(r,8).value=f"=E{r}*F{r}*G{r}"; m.cell(r,8).number_format="#,##0"
    m.cell(r,9,round(x["s_ebitda_unit"],1)).number_format="#,##0.0"
    m.cell(r,9).fill=INF; m.cell(r,9).alignment=CEN; m.cell(r,9).border=BORD
    # monthly EBITDA at UAE realistic sales (from model), only if included
    m.cell(r,10).value=f"=E{r}*{x['mo_ebitda_pot']}"; m.cell(r,10).number_format="#,##0"
    for j in range(1,len(cols)+1): m.cell(r,j).border=BORD
last=first+len(heroes)
# helper columns hidden: electrical flag numeric in col L
for i,(a,qty,elec,inc) in enumerate(heroes):
    m.cell(first+i,12,elec)
m.column_dimensions["L"].hidden=True
# summary
b=last+1
def put(r,label,formula,fmt="#,##0",bold=False,color=None):
    m.cell(r,2,label).alignment=LEFT
    if bold: m.cell(r,2).font=BOLD
    c=m.cell(r,8); c.value=formula; c.number_format=fmt; c.alignment=CEN; c.border=BORD
    if bold: c.font=BOLD
    if color: c.font=Font(bold=True,color=color)
m.cell(b,2,"ИТОГ ЛИН-ЗАПУСКА (по включённым SKU)").font=WHITE; m.cell(b,2).fill=SUB
put(b+1,"Включено SKU",f"=SUM(E{first}:E{last-1})",bold=True)
put(b+2,"Закупка (оборотный капитал), AED",f"=SUM(H{first}:H{last-1})",bold=True)
put(b+3,"Бюджет запуска, AED",
    f"=(SUM(E{first}:E{last-1})*(PHOTO+SAMPLES+PPC_SKU)+LOGI+IF(SUMPRODUCT(E{first}:E{last-1},L{first}:L{last-1})>0,CERT,0))*(1+RESERVE)",bold=True)
put(b+4,"ИТОГО стартовые инвестиции, AED",f"=H{b+2}+H{b+3}",bold=True,color="C00000")
put(b+5,"в USD",f"=H{b+4}/USD_AED","$#,##0")
put(b+6,"EBITDA/мес (полный разгон), AED",f"=SUM(J{first}:J{last-1})",bold=True)
put(b+7,"Прибыль/мес после налога 9%, AED",f"=H{b+6}*(1-CIT)")
put(b+8,"Окупаемость, мес",f"=H{b+4}/(H{b+6}*(1-CIT))","0.0",bold=True,color="1F6E28")
put(b+9,"ROI за 12 мес",f"=H{b+6}*(1-CIT)*12/H{b+4}","0%")
for r in range(b,b+10):
    for j in (2,8): m.cell(r,j).border=BORD

wb.calculation=CalcProperties(fullCalcOnLoad=True)
out=f"{BASE}/deliverables/UAE_Lean_Launch_Model_2026-07.xlsx"
wb.save(out); print("saved",out)
