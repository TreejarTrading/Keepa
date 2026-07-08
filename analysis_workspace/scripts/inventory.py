#!/usr/bin/env python3
"""Этап 1: Инвентаризация всех файлов (xlsx/docx) для анализа Amazon EU."""
import os, sys, hashlib, json
from pathlib import Path
import openpyxl
import docx as docxlib

UPLOADS = Path("/root/.claude/uploads/2e185e71-9c92-5dec-8bf2-5342d7fe62bb")
REPO = Path("/home/user/Keepa")
OUT = REPO / "analysis_workspace" / "01_inventory"
(OUT / "docx_dumps").mkdir(parents=True, exist_ok=True)
(OUT / "xlsx_dumps").mkdir(parents=True, exist_ok=True)

def md5(p, chunk=1 << 20):
    h = hashlib.md5()
    with open(p, "rb") as f:
        while b := f.read(chunk):
            h.update(b)
    return h.hexdigest()

def dump_docx(p: Path, out: Path):
    d = docxlib.Document(str(p))
    lines = []
    for para in d.paragraphs:
        t = para.text.strip()
        if t:
            style = para.style.name if para.style else ""
            prefix = "## " if style.startswith("Heading") else ""
            lines.append(prefix + t)
    for ti, table in enumerate(d.tables):
        lines.append(f"\n=== TABLE {ti+1} ({len(table.rows)}x{len(table.columns)}) ===")
        for row in table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            lines.append(" | ".join(cells))
    out.write_text("\n".join(lines), encoding="utf-8")
    return len(lines)

def dump_xlsx(p: Path, out: Path, max_rows=25):
    try:
        wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
    except Exception as e:
        out.write_text(f"ERROR: {e}")
        return {}
    info = {}
    lines = []
    for ws in wb.worksheets:
        dims = f"{ws.max_row or '?'}x{ws.max_column or '?'}"
        info[ws.title] = dims
        lines.append(f"\n===== SHEET: {ws.title} ({dims}) =====")
        for ri, row in enumerate(ws.iter_rows(values_only=True)):
            if ri >= max_rows:
                more = (ws.max_row - max_rows) if ws.max_row else "?"
                lines.append(f"... ({more} more rows)")
                break
            vals = ["" if v is None else str(v)[:80] for v in row]
            # обрезаем пустые хвосты
            while vals and vals[-1] == "":
                vals.pop()
            lines.append(" | ".join(vals)[:1500])
    wb.close()
    out.write_text("\n".join(lines), encoding="utf-8")
    return info

# Собираем все файлы: uploads + репозиторий
files = []
for p in sorted(UPLOADS.glob("*")):
    files.append(p)
for pattern in ["*.xlsx", "*.pdf", "Products/*.xlsx", "Products/*.csv"]:
    for p in sorted(REPO.glob(pattern)):
        files.append(p)

seen_md5 = {}
inventory = []
for p in files:
    h = md5(p)
    dup_of = seen_md5.get(h)
    entry = {"file": p.name, "path": str(p), "size": p.stat().st_size,
             "md5": h[:8], "duplicate_of": dup_of}
    if not dup_of:
        seen_md5[h] = p.name
        stem = p.stem.split("-", 1)[-1] if p.parent == UPLOADS else p.stem
        stem = stem.replace(" ", "_")[:80]
        if p.suffix == ".docx":
            n = dump_docx(p, OUT / "docx_dumps" / f"{stem}.txt")
            entry["dump"] = f"docx_dumps/{stem}.txt ({n} lines)"
        elif p.suffix == ".xlsx":
            sheets = dump_xlsx(p, OUT / "xlsx_dumps" / f"{stem}.txt")
            entry["sheets"] = sheets
            entry["dump"] = f"xlsx_dumps/{stem}.txt"
    inventory.append(entry)

with open(OUT / "file_inventory.json", "w", encoding="utf-8") as f:
    json.dump(inventory, f, ensure_ascii=False, indent=1)

# Краткая сводка
uniq = [e for e in inventory if not e["duplicate_of"]]
dups = [e for e in inventory if e["duplicate_of"]]
print(f"Всего файлов: {len(inventory)}, уникальных: {len(uniq)}, дубликатов: {len(dups)}")
for e in uniq:
    sheets = ", ".join(f"{k}({v})" for k, v in e.get("sheets", {}).items()) if e.get("sheets") else ""
    print(f"  {e['file'][:70]:<72} {sheets[:100]}")
