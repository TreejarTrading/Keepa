"""Build the launch shortlist: US demand proof + DE actuals + UAE estimate.

Launch markets are DE and UAE (production ~30d in China, sea freight 40-45d
to Jebel Ali / Hamburg). Keepa does not cover Amazon.ae, so the UAE side is
modelled from the US price (x0.95) with Amazon.ae fee assumptions; the DE
side uses real amazon.de prices/demand fetched by de_check.py.

The recommended purchase ceiling («закупка не выше») is the MINIMUM of the
DE and UAE max-COGS at the 25% net-margin target, i.e. the same batch stays
profitable in both destination markets.

Usage:
    uv run python sourcing/shortlist.py
"""

from __future__ import annotations

import csv
import json
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
import economics  # noqa: E402

AE_PRICE_FACTOR = 0.95  # amazon.ae price proxy from the US price


def ae_scenario(price_usd: float, weight_g: float | None) -> dict:
    """Unit economics for Amazon.ae given an expected USD price."""
    net_rev = price_usd * AE_PRICE_FACTOR * (1 - economics.PROMO_DISCOUNT)
    referral = net_rev * economics.DEFAULT_REFERRAL
    w = weight_g or economics.DEFAULT_WEIGHT_G
    fba = next(fee for max_w, fee in economics.FBA_TIERS["AE"] if w <= max_w)
    ppc = net_rev * economics.PPC_SHARE["AE"]
    returns = net_rev * economics.RETURNS_SHARE
    storage_misc = economics.STORAGE_USD["AE"] + economics.MISC_USD["AE"]
    contribution = net_rev - referral - fba - ppc - returns - storage_misc
    freight = max(economics.FREIGHT_MIN_USD, w / 1000 * economics.FREIGHT_USD_PER_KG["AE"])
    duty = economics.DUTY["AE"]
    f25 = (contribution - net_rev * economics.TARGET_MARGIN_RECOMMENDED - freight) / (1 + duty)
    return {
        "ae_price_usd": round(price_usd * AE_PRICE_FACTOR, 2),
        "ae_contribution": round(contribution, 2),
        "ae_max_cogs_25": round(f25, 2),
        "ae_net_unit_25": round(net_rev * economics.TARGET_MARGIN_RECOMMENDED, 2),
    }


COLS = [
    ("asin", "ASIN"), ("title", "Название"), ("category", "Категория"),
    ("us_price", "US цена, $"), ("us_sold", "US продажи/мес"),
    ("us_rating", "Рейтинг"), ("us_reviews", "Отзывы"),
    ("listed_since", "Дата листинга"),
    ("de_listed", "Есть на DE"), ("de_price_eur", "DE цена, €"),
    ("de_sold", "DE продажи/мес"), ("de_offers", "DE офферы"),
    ("de_max_cogs_25", "DE закупка ≤ (25%), $"),
    ("ae_price_usd", "UAE цена (оценка), $"),
    ("ae_max_cogs_25", "UAE закупка ≤ (25%), $"),
    ("target_cogs", "ЗАКУПКА НЕ ВЫШЕ (DE+UAE), $"),
    ("est_net_unit", "Net profit/шт @25%, $"),
    ("est_monthly_net", "Оценка net/мес DE+UAE, $"),
    ("verdict", "Вердикт"), ("rationale", "Обоснование"),
    ("alibaba_url", "Alibaba поиск"), ("url", "Amazon US"),
]


def main() -> None:
    data = Path("sourcing/data")
    us_records = {
        r["asin"]: r
        for r in json.loads((data / "candidates_US.json").read_text(encoding="utf-8"))
        if economics.listed_since_ok(r)
    }
    de_records = {}
    p = data / "candidates_DE.json"
    if p.exists():
        de_records = {r["asin"]: r for r in json.loads(p.read_text(encoding="utf-8"))}
    checked: set[str] = set(de_records)
    p_missing = data / "de_missing.json"
    if p_missing.exists():
        checked |= set(json.loads(p_missing.read_text(encoding="utf-8")))

    rows = []
    for asin in checked:
        rec = us_records.get(asin)
        if not rec:
            continue
        us = economics.compute(rec, "US")
        if not us:
            continue
        de_rec = de_records.get(asin)
        de = economics.compute(de_rec, "DE") if de_rec else None
        ae = ae_scenario(us["price_usd"], rec.get("package_weight_g"))

        cogs_candidates = [ae["ae_max_cogs_25"]]
        if de:
            cogs_candidates.append(de["max_factory_cogs_25"])
        target_cogs = round(min(cogs_candidates), 2)

        de_sold = de["monthly_sold"] if de else None
        est_monthly_units = (de_sold or 0) + max((us["monthly_sold"] or 0) * 0.10, 30)
        net_unit = min(
            [ae["ae_net_unit_25"]] + ([de["net_profit_unit_25"]] if de else [])
        )
        verdict = us["verdict"]
        notes = []
        if de:
            notes.append(f"DE: есть листинг, ≈{de_sold or '?'} шт/мес, {de['offers'] or '?'} офферов")
            if (de_sold or 0) >= 100:
                notes.append("спрос на DE подтверждён")
        else:
            notes.append("на DE не продаётся — свободная ниша, спрос проверять запуском/PPC")
        if target_cogs < 2:
            verdict = "SKIP"
            notes.append("экономика DE/UAE не сходится")
        notes.append(us["rationale"])
        rows.append({
            "asin": asin,
            "title": us["title"],
            "category": us["category"],
            "us_price": us["price_usd"],
            "us_sold": us["monthly_sold"],
            "us_rating": us["rating"],
            "us_reviews": us["reviews"],
            "listed_since": us["listed_since"],
            "de_listed": "да" if de else "нет",
            "de_price_eur": de["price_local"] if de else None,
            "de_sold": de_sold,
            "de_offers": de["offers"] if de else None,
            "de_max_cogs_25": de["max_factory_cogs_25"] if de else None,
            **ae,
            "target_cogs": target_cogs,
            "est_net_unit": round(net_unit, 2),
            "est_monthly_net": round(est_monthly_units * net_unit, 0),
            "verdict": verdict,
            "rationale": "; ".join(notes),
            "alibaba_url": us["alibaba_url"],
            "url": us["url"],
        })

    rows.sort(key=lambda r: (r["verdict"] == "BUY", r["est_monthly_net"] or 0), reverse=True)

    stamp = date.today().strftime("%Y%m%d")
    out = Path("Products") / f"launch_shortlist_DE_UAE_{stamp}.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Shortlist DE+UAE"
    header_fill = PatternFill("solid", fgColor="1F4E79")
    for c, (_, label) in enumerate(COLS, 1):
        cell = ws.cell(row=1, column=c, value=label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    verdict_fill = {
        "BUY": PatternFill("solid", fgColor="C6EFCE"),
        "WATCH": PatternFill("solid", fgColor="FFEB9C"),
        "SKIP": PatternFill("solid", fgColor="FFC7CE"),
    }
    for r, row in enumerate(rows, 2):
        for c, (key, _) in enumerate(COLS, 1):
            val = row.get(key)
            cell = ws.cell(row=r, column=c)
            if key == "asin" and row.get("url"):
                cell.value, cell.hyperlink = val, row["url"]
                cell.font = Font(color="0563C1", underline="single")
            elif key in ("alibaba_url", "url") and val:
                cell.value, cell.hyperlink = ("Alibaba" if key == "alibaba_url" else "Amazon"), val
                cell.font = Font(color="0563C1", underline="single")
            else:
                cell.value = val
            if key == "verdict" and val in verdict_fill:
                cell.fill = verdict_fill[val]
    widths = {"Название": 55, "Обоснование": 70}
    for c, (_, label) in enumerate(COLS, 1):
        ws.column_dimensions[get_column_letter(c)].width = widths.get(label, 15)
    ws.freeze_panes = "C2"
    wb.save(out)

    with (Path("Products") / f"launch_shortlist_DE_UAE_{stamp}.csv").open(
        "w", newline="", encoding="utf-8"
    ) as fh:
        w = csv.DictWriter(fh, fieldnames=[k for k, _ in COLS])
        w.writeheader()
        for r in rows:
            w.writerow({k: r.get(k) for k, _ in COLS})
    print(f"shortlist rows={len(rows)} BUY={sum(1 for r in rows if r['verdict']=='BUY')} -> {out}")


if __name__ == "__main__":
    main()
