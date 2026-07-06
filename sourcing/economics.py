"""Unit economics + ranking for sourcing candidates fetched by fetch_candidates.py.

For every candidate the model computes, per unit and in USD:
    effective revenue after the committed >=10% coupon/deal discount,
    Amazon referral fee, FBA fulfilment fee (Keepa fbaFees or weight-tier
    estimate), storage, PPC (SP+SB launch budget), returns, misc overhead,
    then solves for the MAXIMUM factory price (COGS ceiling) that still
    leaves the target net margin (20% floor / 25% recommended), given
    sea freight ($/kg) and import duty per destination market.

Outputs one master XLSX + CSV into Products/ with clickable ASIN and
Alibaba supplier-search links, plus BUY/WATCH/SKIP verdicts with a Russian
rationale for the purchasing team.

Usage:
    uv run python sourcing/economics.py [--data sourcing/data] [--out Products]
"""

from __future__ import annotations

import argparse
import csv
import json
import re
import urllib.parse
from datetime import date, datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# ----------------------------------------------------------------------------
# Model assumptions (documented in the report; edit here to re-run scenarios)
# ----------------------------------------------------------------------------
FX_TO_USD = {"US": 1.0, "GB": 1.28, "UK": 1.28, "DE": 1.10, "AE": 1 / 3.6725}

PROMO_DISCOUNT = 0.10          # committed minimum coupons/deals/Prime discount
PPC_SHARE = {"US": 0.15, "UK": 0.15, "GB": 0.15, "DE": 0.15, "AE": 0.12}
RETURNS_SHARE = 0.03           # refunds + refurb loss, share of net revenue
STORAGE_USD = {"US": 0.50, "UK": 0.45, "GB": 0.45, "DE": 0.50, "AE": 0.40}
MISC_USD = {"US": 0.35, "UK": 0.20, "GB": 0.20, "DE": 0.20, "AE": 0.20}
DEFAULT_REFERRAL = 0.15

# landed-cost side
FREIGHT_USD_PER_KG = {"US": 1.10, "UK": 1.00, "GB": 1.00, "DE": 1.00, "AE": 0.80}
FREIGHT_MIN_USD = 0.25
DUTY = {"US": 0.35, "UK": 0.04, "GB": 0.04, "DE": 0.04, "AE": 0.05}  # import VAT recoverable, excluded
DEFAULT_WEIGHT_G = 800

# The brief: only listings actually created from Feb 2026. The finder filters
# on trackingSince, which lets old-but-newly-tracked products through, so we
# re-check the true listedSince here (Keepa minutes).
LISTED_SINCE_MIN = int(datetime(2026, 2, 1, tzinfo=timezone.utc).timestamp() // 60) - 21_564_000


def listed_since_ok(record: dict) -> bool:
    ls = record.get("listed_since_keepa")
    if ls is None or ls <= 0:  # unknown -> trust the trackingSince filter
        return True
    return ls >= LISTED_SINCE_MIN


def keepa_minutes_to_date(v) -> str | None:
    try:
        v = int(v)
    except (TypeError, ValueError):
        return None
    if v <= 0:
        return None
    return datetime.fromtimestamp((v + 21_564_000) * 60, tz=timezone.utc).date().isoformat()

TARGET_MARGIN_FLOOR = 0.20
TARGET_MARGIN_RECOMMENDED = 0.25

# FBA fulfilment fee estimate by package weight, used when Keepa fbaFees is
# missing (USD; blended standard-size 2026-level tiers per market).
FBA_TIERS = {  # (max_weight_g, fee_usd)
    "US": [(113, 3.30), (227, 3.45), (340, 3.75), (454, 4.20), (680, 4.70),
           (907, 5.15), (1134, 5.50), (1361, 5.90), (2268, 6.55), (4536, 8.10),
           (9072, 10.60), (1e9, 14.00)],
    "UK": [(150, 2.30), (400, 3.00), (900, 3.65), (1400, 4.30), (2900, 5.20),
           (5900, 6.80), (1e9, 9.50)],
    "DE": [(150, 2.60), (400, 3.30), (900, 3.95), (1400, 4.60), (2900, 5.60),
           (5900, 7.30), (1e9, 10.20)],
    "AE": [(250, 2.60), (500, 2.90), (1000, 3.30), (2000, 3.90), (5000, 5.10),
           (1e9, 7.60)],
}
FBA_TIERS["GB"] = FBA_TIERS["UK"]

# Established / gated / Amazon-owned brands: the listing proves the niche,
# but we sell a private-label analogue, so these are capped at WATCH.
KNOWN_BRANDS = {
    "amazon", "amazon basics", "blink", "ring", "echo", "kindle", "levoit",
    "shark", "ninja", "dyson", "kitchenaid", "cuisinart", "oxo", "rubbermaid",
    "sterilite", "3m", "scotch", "command", "dewalt", "makita", "bosch",
    "black+decker", "blackolt decker", "stanley", "milwaukee", "ryobi", "hp",
    "brother", "sharpie", "post-it", "bissell", "tineco", "irobot", "yeti",
    "hydrojug", "owala", "contigo", "keter", "gorilla", "wd-40", "x-sense",
    "govee", "tp-link", "anker", "eufy", "philips", "braun", "leifheit",
    "kidde", "myq", "chamberlain", "first alert", "ge", "honeywell",
}

STOPWORDS = {
    "with", "for", "and", "the", "of", "in", "to", "set", "pack", "pcs", "new",
    "large", "small", "black", "white", "grey", "gray", "blue", "green", "red",
    "inch", "cm", "mm", "ft", "count", "x",
}


def fba_fee_usd(record: dict, market: str) -> tuple[float, str]:
    fees = record.get("fba_fees") or {}
    ppf = fees.get("pickAndPackFee")
    if ppf and ppf > 0:
        return round(ppf / 100 * FX_TO_USD.get(market, 1.0), 2), "keepa"
    w = record.get("package_weight_g") or DEFAULT_WEIGHT_G
    if w <= 0:
        w = DEFAULT_WEIGHT_G
    for max_w, fee in FBA_TIERS.get(market, FBA_TIERS["US"]):
        if w <= max_w:
            return fee, "estimate"
    return FBA_TIERS[market][-1][1], "estimate"


def pick_price(record: dict) -> float | None:
    """Conservative sell price: min(avg90, current) of the first present series."""
    pricing = (record.get("metrics") or {}).get("pricing") or {}
    for key in ("buy_box", "new", "amazon", "new_fba"):
        s = pricing.get(key) or {}
        vals = [float(s[f]) for f in ("avg", "current") if s.get(f) and s[f] > 0]
        if vals:
            return min(vals)
    return None


def is_known_brand(brand: str | None) -> bool:
    if not brand:
        return False
    b = brand.strip().lower()
    return b in KNOWN_BRANDS or any(b.startswith(k + " ") for k in KNOWN_BRANDS)


def alibaba_keywords(title: str | None, brand: str | None) -> str:
    if not title:
        return ""
    t = title.lower()
    if brand:
        t = t.replace(brand.lower(), " ")
    t = re.sub(r"[^a-z ]+", " ", t)
    words = [w for w in t.split() if len(w) > 2 and w not in STOPWORDS]
    seen, out = set(), []
    for w in words:
        if w not in seen:
            seen.add(w)
            out.append(w)
        if len(out) == 5:
            break
    return " ".join(out)


def compute(record: dict, market: str) -> dict | None:
    fx = FX_TO_USD.get(market, 1.0)
    price_local = pick_price(record)
    if not price_local:
        return None
    price = price_local * fx
    net_rev = price * (1 - PROMO_DISCOUNT)

    ref_pct = record.get("referral_fee_percent")
    ref_pct = (ref_pct / 100 if ref_pct and ref_pct > 1 else ref_pct) or DEFAULT_REFERRAL
    referral = net_rev * ref_pct
    fba, fba_src = fba_fee_usd(record, market)
    storage = STORAGE_USD.get(market, 0.5)
    misc = MISC_USD.get(market, 0.3)
    ppc = net_rev * PPC_SHARE.get(market, 0.15)
    returns = net_rev * RETURNS_SHARE

    contribution = net_rev - referral - fba - storage - misc - ppc - returns

    w_kg = (record.get("package_weight_g") or DEFAULT_WEIGHT_G) / 1000
    freight = max(FREIGHT_MIN_USD, w_kg * FREIGHT_USD_PER_KG.get(market, 1.0))
    duty = DUTY.get(market, 0.05)

    def factory_max(margin: float) -> float:
        landed_max = contribution - net_rev * margin
        return (landed_max - freight) / (1 + duty)

    f20 = factory_max(TARGET_MARGIN_FLOOR)
    f25 = factory_max(TARGET_MARGIN_RECOMMENDED)
    landed_25 = f25 * (1 + duty) + freight if f25 > 0 else None
    net_profit_25 = net_rev * TARGET_MARGIN_RECOMMENDED
    roi_25 = (net_profit_25 / landed_25) if landed_25 and landed_25 > 0 else None

    m = record.get("metrics") or {}
    monthly = record.get("monthly_sold") or (m.get("demand") or {}).get("monthly_sold_estimate") or 0
    offers = ((m.get("competition") or {}).get("offer_count") or {}).get("current")
    bb_amazon = bool((m.get("competition") or {}).get("buy_box_is_amazon"))
    rating = (m.get("reviews") or {}).get("rating_current")
    reviews = (m.get("reviews") or {}).get("review_count_current")
    vol = ((m.get("pricing") or {}).get("new") or {}).get("volatility")

    score = monthly * net_rev * TARGET_MARGIN_RECOMMENDED
    if offers and offers > 12:
        score *= 0.85
    if bb_amazon:
        score *= 0.70
    if vol and vol > 0.4:
        score *= 0.90
    if w_kg > 2:
        score *= 0.90
    if (rating or 0) >= 4.5 and (reviews or 0) >= 50:
        score *= 1.05

    # verdict
    pros, cons = [], []
    if monthly >= 200:
        pros.append(f"продажи ≈{int(monthly)}/мес")
    if f25 >= 3:
        pros.append(f"вход при закупке ≤${f25:.2f} даёт маржу 25%")
    if not bb_amazon:
        pros.append("Buy Box не у Amazon")
    else:
        cons.append("Buy Box у Amazon")
    if offers is not None and offers <= 15:
        pros.append(f"конкуренция умеренная ({int(offers)} офферов)")
    elif offers is not None and offers > 20:
        cons.append(f"много офферов ({int(offers)})")
    if vol is not None and vol > 0.6:
        cons.append(f"волатильность цены {vol}")
    if f20 < 1.5:
        cons.append(f"экономика не сходится (закупка ≤${max(f20,0):.2f} при марже 20%)")
    if w_kg > 3:
        cons.append(f"тяжёлый товар {w_kg:.1f} кг")

    brand_risk = is_known_brand(record.get("brand"))
    if brand_risk:
        cons.append("известный бренд — листинг = референс ниши, продаём PL-аналог")
    if f20 < 1.5 or (bb_amazon and (offers or 0) > 20):
        verdict = "SKIP"
    elif brand_risk:
        verdict = "WATCH"
    elif f25 >= 3 and monthly >= 200 and not bb_amazon and (offers or 99) <= 15:
        verdict = "BUY"
    else:
        verdict = "WATCH"
    confidence = "high" if monthly >= 300 and reviews and reviews >= 30 else "medium"

    kw = alibaba_keywords(record.get("title"), record.get("brand"))
    return {
        "market": market,
        "asin": record.get("asin"),
        "title": record.get("title"),
        "brand": record.get("brand"),
        "category": " > ".join((record.get("category_tree") or [])[:2]),
        "found_by": ",".join(record.get("found_by") or []),
        "url": record.get("url"),
        "alibaba_url": "https://www.alibaba.com/trade/search?SearchText=" + urllib.parse.quote(kw) if kw else None,
        "price_local": round(price_local, 2),
        "price_usd": round(price, 2),
        "net_rev_after_promo": round(net_rev, 2),
        "referral_pct": round(ref_pct * 100, 1),
        "referral_usd": round(referral, 2),
        "fba_fee_usd": round(fba, 2),
        "fba_fee_source": fba_src,
        "ppc_usd": round(ppc, 2),
        "returns_usd": round(returns, 2),
        "storage_misc_usd": round(storage + misc, 2),
        "contribution_usd": round(contribution, 2),
        "freight_usd": round(freight, 2),
        "duty_pct": round(duty * 100, 1),
        "max_factory_cogs_20": round(f20, 2),
        "max_factory_cogs_25": round(f25, 2),
        "landed_at_25": round(landed_25, 2) if landed_25 else None,
        "net_profit_unit_25": round(net_profit_25, 2),
        "roi_25_pct": round(roi_25 * 100, 1) if roi_25 else None,
        "monthly_sold": int(monthly) if monthly else None,
        "monthly_net_potential_usd": round(score, 0),
        "rating": round(rating, 1) if rating is not None else None,
        "reviews": int(reviews) if reviews else None,
        "offers": int(offers) if offers is not None else None,
        "bb_amazon": bb_amazon,
        "volatility": vol,
        "weight_g": record.get("package_weight_g"),
        "listed_since": keepa_minutes_to_date(record.get("listed_since_keepa"))
        or keepa_minutes_to_date(record.get("tracking_since_keepa")),
        "verdict": verdict,
        "confidence": confidence,
        "rationale": "; ".join(["+ " + p for p in pros] + ["− " + c for c in cons]),
    }


COLUMNS = [
    ("market", "Рынок"), ("asin", "ASIN"), ("title", "Название"),
    ("category", "Категория"), ("price_usd", "Цена, $"),
    ("monthly_sold", "Продажи/мес"), ("rating", "Рейтинг"),
    ("reviews", "Отзывы"), ("offers", "Офферы"), ("bb_amazon", "BB Amazon"),
    ("net_rev_after_promo", "Выручка после промо 10%, $"),
    ("referral_usd", "Referral, $"), ("fba_fee_usd", "FBA, $"),
    ("ppc_usd", "PPC 15%, $"), ("returns_usd", "Возвраты 3%, $"),
    ("storage_misc_usd", "Хранение+проч., $"),
    ("contribution_usd", "Остаток до COGS, $"),
    ("freight_usd", "Фрахт/шт, $"), ("duty_pct", "Пошлина, %"),
    ("max_factory_cogs_25", "Закупка НЕ ВЫШЕ (маржа 25%), $"),
    ("max_factory_cogs_20", "Закупка предел (маржа 20%), $"),
    ("landed_at_25", "Landed @25%, $"), ("net_profit_unit_25", "Net profit/шт @25%, $"),
    ("roi_25_pct", "ROI @25%, %"),
    ("monthly_net_potential_usd", "Потенциал net/мес, $"),
    ("listed_since", "Дата листинга"),
    ("verdict", "Вердикт"), ("confidence", "Уверенность"),
    ("rationale", "Обоснование"), ("alibaba_url", "Alibaba поиск"),
    ("fba_fee_source", "FBA fee источник"), ("volatility", "Волатильность"),
    ("weight_g", "Вес, г"), ("brand", "Бренд"), ("found_by", "Найден поиском"),
]


def write_xlsx(rows: list[dict], path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Master"
    _fill_sheet(ws, rows)
    ws_buy = wb.create_sheet("TOP BUY")
    _fill_sheet(ws_buy, [r for r in rows if r["verdict"] == "BUY"])
    ws_a = wb.create_sheet("Допущения")
    assumptions = [
        ("Промо/купоны/deals (минимум)", f"{PROMO_DISCOUNT:.0%} от цены на все продажи"),
        ("PPC SP+SB (год запуска)", "15% выручки (12% UAE)"),
        ("Возвраты", f"{RETURNS_SHARE:.0%} выручки"),
        ("Хранение FBA + прочее", "US $0.85/шт; UK/DE/AE $0.6-0.7/шт"),
        ("Referral fee", "из Keepa; иначе 15%"),
        ("FBA fee", "из Keepa fbaFees; иначе оценка по весу"),
        ("Фрахт (море, LCL, 40-45 дней)", "US $1.1/кг; DE/UK $1.0/кг; UAE $0.8/кг"),
        ("Пошлина", "US 35% (Китай); UK/DE 4%; UAE 5%. Импортный VAT возвратный — не в модели"),
        ("Целевая маржа", "рекомендованная 25% net; минимально допустимая 20%"),
        ("Таймлайн", "производство ~30 дн + море 40-45 дн + приёмка ~10 дн ≈ 80-85 дней до старта продаж"),
        ("FX", "GBP=1.28$, EUR=1.10$, AED=0.2723$"),
        ("Дата расчёта", str(date.today())),
    ]
    for i, (k, v) in enumerate(assumptions, 1):
        ws_a.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws_a.cell(row=i, column=2, value=v)
    ws_a.column_dimensions["A"].width = 40
    ws_a.column_dimensions["B"].width = 80
    wb.save(path)


def _fill_sheet(ws, rows: list[dict]) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E79")
    for c, (_, label) in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=c, value=label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    verdict_fill = {
        "BUY": PatternFill("solid", fgColor="C6EFCE"),
        "WATCH": PatternFill("solid", fgColor="FFEB9C"),
        "SKIP": PatternFill("solid", fgColor="FFC7CE"),
    }
    for r, row in enumerate(rows, 2):
        for c, (key, _) in enumerate(COLUMNS, 1):
            val = row.get(key)
            cell = ws.cell(row=r, column=c)
            if key == "asin" and row.get("url"):
                cell.value = val
                cell.hyperlink = row["url"]
                cell.font = Font(color="0563C1", underline="single")
            elif key == "alibaba_url" and val:
                cell.value = "Alibaba"
                cell.hyperlink = val
                cell.font = Font(color="0563C1", underline="single")
            else:
                cell.value = val
            if key == "verdict" and val in verdict_fill:
                cell.fill = verdict_fill[val]
    widths = {"Название": 55, "Обоснование": 60, "Категория": 30}
    for c, (_, label) in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(c)].width = widths.get(label, 14)
    ws.freeze_panes = "C2"


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--data", default="sourcing/data")
    ap.add_argument("--out", default="Products")
    args = ap.parse_args()
    data_dir, out_dir = Path(args.data), Path(args.out)
    out_dir.mkdir(exist_ok=True)

    rows: list[dict] = []
    for f in sorted(data_dir.glob("candidates_*.json")):
        market = f.stem.split("_")[-1]
        records = json.loads(f.read_text(encoding="utf-8"))
        skipped_old = 0
        for rec in records:
            if not listed_since_ok(rec):
                skipped_old += 1
                continue
            row = compute(rec, market)
            if row:
                rows.append(row)
        if skipped_old:
            print(f"{f.name}: skipped {skipped_old} listings older than 2026-02")
    rows.sort(key=lambda r: r.get("monthly_net_potential_usd") or 0, reverse=True)

    stamp = date.today().strftime("%Y%m%d")
    xlsx = out_dir / f"sourcing_master_{stamp}.xlsx"
    write_xlsx(rows, xlsx)
    with (out_dir / f"sourcing_master_{stamp}.csv").open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=[k for k, _ in COLUMNS] + ["url"])
        w.writeheader()
        for r in rows:
            w.writerow({k: r.get(k) for k in list(dict(COLUMNS)) + ["url"]})
    counts = {v: sum(1 for r in rows if r["verdict"] == v) for v in ("BUY", "WATCH", "SKIP")}
    print(f"rows={len(rows)} verdicts={counts} -> {xlsx}")


if __name__ == "__main__":
    main()
