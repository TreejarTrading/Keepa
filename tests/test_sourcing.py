"""Offline tests for the China-sourcing layer (no Keepa API)."""

from __future__ import annotations

import pytest
from openpyxl import load_workbook


# --- MOQ price tiers (problem #1) -------------------------------------------

def test_normalize_tiers_infers_max_qty():
    from keepa_mcp import sourcing

    tiers = sourcing.normalize_tiers([
        {"min_qty": 1, "max_qty": 100, "unit_price": 9.91},
        {"min_qty": 1000, "unit_price": 9.53},  # open tier, no max
        {"min_qty": 101, "max_qty": 999, "unit_price": 9.78},
    ])
    assert [t["min_qty"] for t in tiers] == [1, 101, 1000]
    assert tiers[0]["max_qty"] == 100
    assert tiers[2]["max_qty"] is None  # last tier stays open


def test_pick_tier_matches_order_quantity():
    from keepa_mcp import sourcing

    tiers = [
        {"min_qty": 1, "max_qty": 100, "unit_price": 9.91},
        {"min_qty": 101, "max_qty": 999, "unit_price": 9.78},
        {"min_qty": 1000, "max_qty": None, "unit_price": 9.53},
    ]
    # 100 units -> first tier (NOT the cheapest >=1000 tier) — the core fix.
    assert sourcing.pick_tier(tiers, 100)["unit_price"] == 9.91
    assert sourcing.pick_tier(tiers, 500)["unit_price"] == 9.78
    assert sourcing.pick_tier(tiers, 5000)["unit_price"] == 9.53
    # qty=None falls back to the MOQ (smallest) tier.
    assert sourcing.pick_tier(tiers, None)["unit_price"] == 9.91


# --- currency ---------------------------------------------------------------

def test_to_base_uses_fx_and_flags_unknown():
    from keepa_mcp import sourcing

    val, known = sourcing.to_base(9.91, "EUR", "USD", {"EUR": 1.08})
    assert val == pytest.approx(10.7028)
    assert known is True
    # Missing rate -> assume 1.0 and flag it.
    val2, known2 = sourcing.to_base(100, "CNY", "USD", {})
    assert val2 == 100 and known2 is False


# --- referral fee by category (problem: not flat 15%) -----------------------

def test_referral_pct_by_category():
    from keepa_mcp import fees

    assert fees.estimate_referral_pct(["Electronics", "Headphones"]) == 0.08
    assert fees.estimate_referral_pct(["Clothing, Shoes & Jewelry", "Men", "Shirts"]) == 0.17
    assert fees.estimate_referral_pct(["Home & Kitchen"]) == 0.15
    # Price-tiered
    assert fees.estimate_referral_pct(["Jewelry"], price=100) == 0.20
    assert fees.estimate_referral_pct(["Jewelry"], price=500) == 0.05
    assert fees.estimate_referral_pct(["Furniture"], price=150) == 0.15
    assert fees.estimate_referral_pct(["Furniture"], price=900) == 0.10
    # Unknown -> default
    assert fees.estimate_referral_pct([]) == 0.15


def test_analysis_extracts_keepa_fees():
    """Keepa's real referral % + FBA pick&pack fee land in the record."""
    from keepa_mcp import analysis

    product = {
        "asin": "B0FEE", "title": "Thing", "categoryTree": [{"name": "Electronics"}],
        "referralFeePercent": 8,                 # Keepa real value (percent)
        "fbaFees": {"pickAndPackFee": 538},      # cents -> $5.38
        "data": {"NEW": [40.0], "NEW_time": [0]},
    }
    rec = analysis.build_record(product, stats_days=90)
    f = rec["metrics"]["amazon_fees"]
    assert f["referral_pct"] == 0.08 and f["referral_pct_source"] == "keepa"
    assert f["fba_fee"] == 5.38
    # referral = 40*0.08=3.20; total = 3.20+5.38=8.58; net = 40-8.58=31.42
    assert f["referral_fee"] == pytest.approx(3.20)
    assert f["total_fees"] == pytest.approx(8.58)
    assert f["net_proceeds"] == pytest.approx(31.42)
    assert rec["fba_pick_pack_fee"] == 5.38


def test_analysis_referral_falls_back_to_category():
    """No Keepa referral field -> category estimate (8% for electronics)."""
    from keepa_mcp import analysis

    product = {
        "asin": "B0CAT", "title": "Camera", "categoryTree": [{"name": "Camera & Photo"}],
        "data": {"NEW": [200.0], "NEW_time": [0]},
    }
    f = analysis.build_record(product)["metrics"]["amazon_fees"]
    assert f["referral_pct"] == 0.08
    assert f["referral_pct_source"] == "category"


# --- economics + verdict ----------------------------------------------------

def _sample_item():
    return {
        "asin": "B0EX", "title": "Dual Monitor Stand", "brand": "Wantai",
        "manufacturer": "Putian Wantai Hardware Co., Ltd.",
        "marketplace": "DE", "currency": "EUR",
        "amazon_sell_price": 32.0, "amazon_referral_pct": 0.15,
        "amazon_referral_source": "keepa", "amazon_fba_fee": 4.50,
        "monthly_sold": 400, "weight_g": 2500,
        "url": "https://www.amazon.de/dp/B0EX",
        "suppliers": [{
            "platform": "1688", "supplier_name": "Putian Wantai Hardware",
            "is_manufacturer": True, "matches_amazon_manufacturer": True,
            "match_quality": "точное", "match_basis": "MS008 + фото",
            "model_number": "MS008", "currency": "CNY", "moq": 100,
            "price_tiers": [
                {"min_qty": 1, "max_qty": 100, "unit_price": 48},
                {"min_qty": 101, "max_qty": 999, "unit_price": 45},
                {"min_qty": 1000, "max_qty": None, "unit_price": 42},
            ],
            "url": "https://detail.1688.com/x.html",
        }],
    }


def test_build_plan_economics_and_verdict():
    from keepa_mcp import sourcing

    plan = sourcing.build_plan(
        [_sample_item()], base_currency="USD",
        fx={"EUR": 1.08, "CNY": 0.14}, duty_pct=0.05, freight_per_kg=1.2,
    )
    row = plan["rows"][0]
    assert row["order_qty"] == 100 and row["tier_range"] == "1–100"
    # unit 48*0.14=6.72; freight 2.5*1.2=3; duty (6.72+3)*0.05=0.486; landed≈10.21
    assert row["unit_cost"] == pytest.approx(6.72)
    assert row["freight"] == pytest.approx(3.0)
    assert row["landed_cost"] == pytest.approx(10.206, abs=0.01)
    assert row["margin_pct"] > 0.30 and row["roi_pct"] > 0.60
    assert row["verdict"] == "ЗАКУПАТЬ"
    assert len(row["scenarios"]) == 3  # one per MOQ tier
    # weight-based freight is recorded in assumptions
    assert "кг ×" in row["assumptions"]


def test_verdict_reject_on_analog_or_thin_margin():
    from keepa_mcp import sourcing

    item = _sample_item()
    item["suppliers"][0]["match_quality"] = "аналог"
    row = sourcing.build_plan([item], fx={"EUR": 1.08, "CNY": 0.14},
                              freight_per_kg=1.2)["rows"][0]
    assert row["verdict"] == "ОТКАЗ"


def test_amazon_item_from_record_carries_manufacturer_and_fees():
    from keepa_mcp import analysis, sourcing

    product = {
        "asin": "B0M", "title": "Cam", "brand": "BrandX", "manufacturer": "FactoryY",
        "categoryTree": [{"name": "Electronics"}],
        "referralFeePercent": 8, "fbaFees": {"pickAndPackFee": 500},
        "monthlySold": 300, "packageWeight": 1800,
        "data": {"BUY_BOX_SHIPPING": [50.0], "BUY_BOX_SHIPPING_time": [0],
                 "NEW": [52.0], "NEW_time": [0]},
    }
    rec = analysis.build_record(product)
    item = sourcing.amazon_item_from_record(rec)
    assert item["manufacturer"] == "FactoryY"
    assert item["amazon_referral_pct"] == 0.08
    assert item["amazon_fba_fee"] == 5.0
    assert item["amazon_sell_price"] == 50.0  # buy box preferred
    assert item["weight_g"] == 1800


# --- report -----------------------------------------------------------------

def test_generate_sourcing_report(tmp_path, monkeypatch):
    monkeypatch.setenv("KEEPA_OUTPUT_DIR", str(tmp_path))
    import importlib
    from keepa_mcp import config, reports, sourcing
    importlib.reload(config)
    importlib.reload(reports)

    plan = sourcing.build_plan(
        [_sample_item()], fx={"EUR": 1.08, "CNY": 0.14}, freight_per_kg=1.2
    )
    path = reports.generate_sourcing_report(plan, report_name="t", query_summary="ctx")
    assert path.exists() and path.suffix == ".xlsx"

    wb = load_workbook(path)
    assert {"Сопоставление", "Сценарии", "Пояснения", "Сводка"}.issubset(wb.sheetnames)
    headers = [c.value for c in wb["Сопоставление"][1]]
    for must in ("Производитель (Amazon)", "Площадка (Китай)", "Степень совпадения",
                 "Ценовой уровень (MOQ)", "Себестоимость (landed)", "Маржа, %", "Вердикт"):
        assert must in headers
    # Пояснения sheet documents every comparison column.
    explained = {r[0].value for r in wb["Пояснения"].iter_rows()}
    assert "Производитель (Amazon)" in explained
    assert "Себестоимость (landed)" in explained


def test_amazon_report_has_manufacturer_and_fee_columns(tmp_path, monkeypatch):
    monkeypatch.setenv("KEEPA_OUTPUT_DIR", str(tmp_path))
    import importlib
    from keepa_mcp import config, reports, analysis
    importlib.reload(config)
    importlib.reload(reports)

    product = {
        "asin": "B0R", "title": "X", "manufacturer": "FactoryZ",
        "categoryTree": [{"name": "Electronics"}],
        "referralFeePercent": 8, "fbaFees": {"pickAndPackFee": 400},
        "data": {"NEW": [30.0], "NEW_time": [0]},
    }
    rec = analysis.build_record(product)
    path = reports.generate_report([rec], report_name="amz")
    headers = [c.value for c in load_workbook(path)["Анализ"][1]]
    for must in ("Manufacturer", "Referral %", "FBA fee", "Amazon fees", "Net after fees"):
        assert must in headers
